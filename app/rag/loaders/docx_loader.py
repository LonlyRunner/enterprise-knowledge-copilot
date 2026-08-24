from pathlib import Path

from docx import Document as DocxDocument

from app.rag.loaders.base import BaseDocumentLoader
from app.rag.models import Document


class DocxDocumentLoader(
    BaseDocumentLoader
):

    def load(
        self,
        file_path: str,
        *,
        document_id: str | None = None,
        tenant_id: str = "default",
    ) -> Document:

        path = Path(file_path)

        doc = DocxDocument(
            file_path
        )

        paragraphs = []

        for paragraph in doc.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(
                    text
                )

        content = "\n\n".join(
            paragraphs
        )

        return Document(
            id=document_id or path.stem,
            tenant_id=tenant_id,
            content=content,
            metadata={
                "source": path.name,
                "file_type": "docx",
                "path": str(path),
            },
        )
